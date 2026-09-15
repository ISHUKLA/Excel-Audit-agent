#!/usr/bin/env python3
"""Recalculate and qualify competition demonstration workbooks.

Workbook authoring remains in ``build_competition_cases.mjs`` and uses
``@oai/artifact-tool``. This finaliser performs the native-engine pass, checks
the saved artifacts through production code, and writes inventories, expected
results and evidence manifests.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from agents.anomaly_detector import detect_anomalies
from agents.parser import parse_workbook
from agents.reconciliation import run_reconciliation
from core.formula_catalogue import SUPPORTED_FUNCTIONS

COMPETITION_DIR = ROOT / "demo" / "final cases"
FUNCTION_PATTERN = re.compile(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(")
BOOLEAN_CALLS = {"TRUE", "FALSE"}

CASE_IDS = {
    7: ("case_7a_ifrs17_clean", "case_7b_ifrs17_missing_cohort"),
    8: ("case_8_solvency_capital",),
    9: ("case_9_life_pricing",),
    10: ("case_10_accounting_context",),
}
CASE_FOLDERS = {
    "case_7a_ifrs17_clean": "case_7_ifrs17",
    "case_7b_ifrs17_missing_cohort": "case_7_ifrs17",
    "case_8_solvency_capital": "case_8_solvency_capital",
    "case_9_life_pricing": "case_9_life_pricing",
    "case_10_accounting_context": "case_10_accounting_context",
}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _formula_inventory(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    inventory: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.startswith("="):
                    continue
                functions = sorted(
                    {
                        name.upper().removeprefix("_XLFN.")
                        for name in FUNCTION_PATTERN.findall(cell.value)
                    }
                    - BOOLEAN_CALLS
                )
                unsupported = sorted(set(functions) - SUPPORTED_FUNCTIONS)
                inventory.append(
                    {
                        "cell_ref": f"{sheet.title}!{cell.coordinate}",
                        "formula": cell.value,
                        "functions": functions,
                        "supported_by_catalogue": not unsupported,
                        "unsupported_functions": unsupported,
                    }
                )
    return inventory


def _formula_cached_values(path: Path) -> dict[str, object]:
    formula_book = load_workbook(path, data_only=False, read_only=False)
    value_book = load_workbook(path, data_only=True, read_only=False)
    values: dict[str, object] = {}
    for formula_sheet in formula_book.worksheets:
        value_sheet = value_book[formula_sheet.title]
        for row in formula_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    values[f"{formula_sheet.title}!{cell.coordinate}"] = value_sheet[
                        cell.coordinate
                    ].value
    return values


def _soffice() -> str:
    executable = (
        os.environ.get("SOFFICE_BIN")
        or shutil.which("soffice")
        or shutil.which("libreoffice")
    )
    if not executable:
        raise RuntimeError(
            "LibreOffice is required to qualify competition workbooks; no final "
            "workbook or evidence manifest was written."
        )
    return executable


def _set_qualified_calc_mode(workbook_path: Path) -> None:
    """Make LibreOffice's completed automatic recalculation explicit.

    This changes calculation metadata only. Formula caches come exclusively
    from LibreOffice and are never manufactured by direct OOXML editing.
    """

    with zipfile.ZipFile(workbook_path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    xml = entries["xl/workbook.xml"].decode("utf-8")
    calc_pr = '<calcPr calcId="191029" calcMode="auto" fullCalcOnLoad="1"/>'
    if "<calcPr" in xml:
        start = xml.index("<calcPr")
        end = xml.index(">", start) + 1
        xml = xml[:start] + calc_pr + xml[end:]
    else:
        xml = xml.replace("</workbook>", calc_pr + "</workbook>")
    entries["xl/workbook.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(workbook_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)


def _recalculate(source: Path, final: Path, executable: str) -> None:
    with tempfile.TemporaryDirectory(prefix=f"competition-{source.stem}-recalc-") as temp:
        subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                temp,
                str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=600,
        )
        produced = Path(temp) / source.name
        if not produced.exists():
            raise RuntimeError(f"LibreOffice did not produce {source.name}")
        final.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(produced, final)
    _set_qualified_calc_mode(final)


def _write_reference_csv(path: Path, rows: list[dict]) -> None:
    columns = [
        "account_number",
        "label",
        "amount",
        "debit_credit",
        "ledger_source",
        "evidence_reference",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row[column] for column in columns})


def _git_evidence() -> tuple[str, bool]:
    commit = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    dirty = bool(
        subprocess.run(
            ["git", "status", "--short"],
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()
    )
    return commit, dirty


def _verify_production_results(final: Path, expected: dict) -> dict[str, object]:
    parsed = parse_workbook(final.read_bytes())
    if parsed.workbook_meta.calc_mode != "automatic":
        raise RuntimeError(f"{final.name}: calculation mode is not automatic")
    formula_errors = [
        ref
        for ref, record in parsed.cells.items()
        if record.formula is not None and record.is_error
    ]
    if formula_errors:
        raise RuntimeError(f"{final.name}: formula errors after recalculation: {formula_errors}")

    findings = detect_anomalies(parsed)
    expected_finding = expected.get("expected_finding")
    if expected_finding is None and findings:
        raise RuntimeError(
            f"{final.name}: unexpected findings: "
            + "; ".join(f"{item.tab}!{item.cell_ref}: {item.description}" for item in findings)
        )
    if expected_finding is not None:
        location = expected_finding["formula_cell"]
        matching = [
            item
            for item in findings
            if f"{item.tab}!{item.cell_ref}" == location
        ]
        if len(matching) != 1:
            raise RuntimeError(
                f"{final.name}: expected one finding at {location}, observed "
                f"{[(item.tab, item.cell_ref, item.description) for item in findings]}"
            )

    output_items = list(expected["summary_outputs"].values())
    result = run_reconciliation(parsed, [item["cell"] for item in output_items])
    internal = [line for line in result.lines if line.check_type == "excel_vs_python"]
    if len(internal) != len(output_items):
        raise RuntimeError(f"{final.name}: did not reconstruct every expected output")
    for item, line in zip(output_items, internal):
        expected_completeness = item.get("expected_completeness", "complete")
        expected_verdict = "pass" if expected_completeness == "complete" else "incomplete"
        if line.completeness != expected_completeness or line.verdict != expected_verdict:
            raise RuntimeError(
                f"{final.name}: {item['cell']} was {line.completeness}/{line.verdict}: "
                f"{line.unsupported_elements}"
            )
        expected_value = item["expected_value"]
        tolerance = item.get("tolerance", 1e-6)
        if line.source_value is None or abs(line.source_value - expected_value) > tolerance:
            raise RuntimeError(
                f"{final.name}: cached {item['cell']}={line.source_value}, expected {expected_value}"
            )
        if expected_completeness == "partial":
            if line.target_value is not None:
                raise RuntimeError(
                    f"{final.name}: partial {item['cell']} invented target {line.target_value}"
                )
            for function in item.get("unsupported_functions", []):
                if not any(function in reason for reason in line.unsupported_elements):
                    raise RuntimeError(
                        f"{final.name}: {item['cell']} did not report unsupported {function}"
                    )
        elif line.target_value is None or abs(line.target_value - expected_value) > tolerance:
            raise RuntimeError(
                f"{final.name}: Python {item['cell']}={line.target_value}, expected {expected_value}"
            )
    return {
        "finding_count": len(findings),
        "findings": [
            {
                "finding_id": item.finding_id,
                "severity": item.severity,
                "cell_ref": f"{item.tab}!{item.cell_ref}",
                "description": item.description,
                "raw_value": item.raw_value,
            }
            for item in findings
        ],
    }


def _case_folder(case_id: str) -> Path:
    return COMPETITION_DIR / CASE_FOLDERS[case_id]


def _finalize_one(
    *,
    metadata_path: Path,
    executable: str,
    engine_version: str,
    recalculated_at: str,
    tool_commit: str,
    worktree_dirty: bool,
) -> None:
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    case_id = metadata["caseId"]
    source = Path(metadata["workbookPath"])
    folder = _case_folder(case_id)
    workbook_path = folder / "workbooks" / source.name
    reference_filename = metadata.get("referenceFilename")
    if reference_filename is None and metadata["referenceLines"]:
        reference_filename = f"{case_id}_reference_figures.csv"
    reference_path = (
        folder / "reference_figures" / reference_filename
        if reference_filename is not None
        else None
    )
    expected_path = folder / "expected_results" / f"{case_id}_expected.json"
    inventory_path = folder / "formula_inventories" / f"{case_id}_formula_inventory.json"
    manifest_path = folder / "manifests" / f"{case_id}_manifest.json"
    for directory in (
        workbook_path.parent,
        expected_path.parent,
        inventory_path.parent,
        manifest_path.parent,
        folder / "docs",
    ):
        directory.mkdir(parents=True, exist_ok=True)
    if reference_path is not None:
        reference_path.parent.mkdir(parents=True, exist_ok=True)

    _recalculate(source, workbook_path, executable)
    if reference_path is not None:
        _write_reference_csv(reference_path, metadata["referenceLines"])
    inventory = _formula_inventory(workbook_path)
    inventory_path.write_text(json.dumps(inventory, indent=2) + "\n", encoding="utf-8")

    expected = metadata["expected"]
    observed = _verify_production_results(workbook_path, expected)
    expected.update(
        {
            "schema_version": 1,
            "case_id": case_id,
            "case_version": metadata["caseVersion"],
            "synthetic_data": True,
            "workbook_file": str(workbook_path.relative_to(ROOT)),
            "workbook_sha256": _sha256(workbook_path),
            "reference_file": (
                str(reference_path.relative_to(ROOT)) if reference_path is not None else None
            ),
            "reference_sha256": _sha256(reference_path) if reference_path is not None else None,
            "reference_control_total": metadata["referenceControlTotal"],
            "formula_inventory_file": str(inventory_path.relative_to(ROOT)),
            "formula_cell_count": len(inventory),
            "supported_formula_count": sum(
                bool(item["supported_by_catalogue"]) for item in inventory
            ),
            "unsupported_formula_count": sum(
                not bool(item["supported_by_catalogue"]) for item in inventory
            ),
            "observed_findings": observed["findings"],
            "tool_git_commit": tool_commit,
            "tool_worktree_dirty_at_generation": worktree_dirty,
        }
    )
    expected_path.write_text(json.dumps(expected, indent=2) + "\n", encoding="utf-8")

    authoritative = [
        {
            "cell": cell,
            "expected_value": next(
                item["expected_value"]
                for item in expected["summary_outputs"].values()
                if item["cell"] == cell
            ),
        }
        for cell in expected["authoritative_output_cells"]
    ]
    manifest = {
        "case_id": case_id,
        "case_version": metadata["caseVersion"],
        "workbook_file": str(workbook_path.relative_to(ROOT)),
        "workbook_sha256": _sha256(workbook_path),
        "reference_file": (
            str(reference_path.relative_to(ROOT)) if reference_path is not None else None
        ),
        "reference_sha256": _sha256(reference_path) if reference_path is not None else None,
        "expected_results_file": str(expected_path.relative_to(ROOT)),
        "expected_results_sha256": _sha256(expected_path),
        "formula_inventory_file": str(inventory_path.relative_to(ROOT)),
        "formula_inventory_sha256": _sha256(inventory_path),
        "recalculation_engine": "LibreOffice",
        "recalculation_engine_version": engine_version,
        "recalculated_at_utc": recalculated_at,
        "tool_git_commit": tool_commit,
        "tool_worktree_dirty_at_generation": worktree_dirty,
        "python_version": platform.python_version(),
        "formula_cell_count": len(inventory),
        "supported_formula_count": expected["supported_formula_count"],
        "unsupported_formula_count": expected["unsupported_formula_count"],
        "authoritative_outputs": authoritative,
        "expected_internal_verdict": expected["expected_internal_verdict"],
        "expected_external_verdict": expected["expected_external_verdict"],
        "expected_gate_block": expected["expected_gate_block"],
        "synthetic_data": True,
    }
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(
        f"{case_id}: {len(inventory)} formulas, SHA-256 {_sha256(workbook_path)}, "
        "production verification passed"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--build-dir", type=Path, required=True)
    parser.add_argument("--cases", nargs="+", type=int, choices=tuple(CASE_IDS), required=True)
    args = parser.parse_args()
    build_dir = args.build_dir.resolve()
    executable = _soffice()
    engine_version = subprocess.run(
        [executable, "--version"],
        check=True,
        capture_output=True,
        text=True,
        timeout=30,
    ).stdout.strip()
    recalculated_at = datetime.now(timezone.utc).isoformat()
    tool_commit, worktree_dirty = _git_evidence()

    for case_number in args.cases:
        for case_id in CASE_IDS[case_number]:
            metadata_path = build_dir / f"{case_id}_builder_metadata.json"
            if not metadata_path.exists():
                raise FileNotFoundError(f"Missing builder metadata: {metadata_path}")
            _finalize_one(
                metadata_path=metadata_path,
                executable=executable,
                engine_version=engine_version,
                recalculated_at=recalculated_at,
                tool_commit=tool_commit,
                worktree_dirty=worktree_dirty,
            )

    print(f"LibreOffice: {engine_version}")
    print("Microsoft Excel compatibility tested: no")


if __name__ == "__main__":
    main()
